"""Fill VITISH'26 Idea Template with Radiology AI project content."""
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt


TEMPLATE = Path(r"c:\Users\abhin\Downloads\VITISH'26_IdeaTemplate.pptx")
OUTPUT = Path(r"c:\Users\abhin\Downloads\VITISH'26_RadiologyAI_Idea.pptx")

# Official VITISH 2026 / SIH Internal Hackathon problem statement
PROBLEM_STATEMENT_ID = "45"
PROBLEM_STATEMENT_TITLE = "AI-powered Radiology Interpretation Assistant"
THEME = "HealthTech"
PS_CATEGORY = "Software"

IDEA_TITLE = "Radiology AI Assistant — Explainable Clinical Decision Support for Chest X-Ray Analysis"

SLIDE2_BULLETS = [
    "Web-based clinical decision-support platform for chest X-ray analysis",
    "Upload PNG/JPG/DICOM → AI predicts 18 pathology classes in ~5 seconds",
    "Grad-CAM heatmaps highlight infection/abnormality regions (explainable AI)",
    "Ranked differential diagnosis with confidence meter and severity scoring",
    "Auto-generated radiology report with findings, impression, and PDF export",
    "Flags low-confidence cases as 'Needs Human Review' — radiologist has final say",
    "Addresses India's radiologist shortage (~1 per 100,000 vs WHO norm 1:10,000)",
    "Cuts rural reporting delays from 3–7 days to seconds for preliminary triage",
    "Beyond black-box detection: full workflow — prediction + evidence + report + PACS",
]

SLIDE3_BULLETS = [
    "AI/ML: PyTorch, DenseNet121 (TorchXRayVision), Grad-CAM (Captum), CLAHE preprocessing",
    "Backend: FastAPI, SQLAlchemy, SQLite, JWT authentication, REST API",
    "Frontend: React, Tailwind CSS, Chart.js dashboard for clinicians",
    "Medical imaging: pydicom, DICOM upload, Orthanc PACS integration (Docker)",
    "Deployment: Docker Compose, Nginx, Redis — production-ready architecture",
    "Pipeline: Upload → Preprocess → Predict → Grad-CAM → Report → Dashboard/PDF",
    "Working prototype with demo samples, patient history, and metrics dashboard",
]

SLIDE4_BULLETS = [
    "Feasible: open-source stack, pre-trained models, Docker deploy on modest hardware",
    "MVP built in 7 days — full-stack with explainability, reports, and PACS hooks",
    "Challenges: clinical validation, dataset bias, regulatory approval for deployment",
    "Risk: false negatives in critical cases — mitigated by confidence thresholds + human review",
    "Risk: model drift across X-ray machines — mitigated by preprocessing + periodic retraining",
    "Strategy: pilot with partner hospital, radiologist-in-the-loop workflow, iterative validation",
]

SLIDE5_BULLETS = [
    "Rural doctors: instant AI-assisted preliminary findings without waiting days",
    "Radiologists: reporting time reduced from ~30 min to ~30 sec for triage",
    "Patients: faster pneumonia/TB detection — 2,500 children die daily from preventable pneumonia",
    "Hospitals: reduced imaging backlog, lower tele-radiology costs",
    "Social: equitable access to diagnostic support in underserved areas",
    "Economic: India medical imaging market ~$2.1B by 2028; tele-radiology growing ~23%/year",
]

SLIDE6_BULLETS = [
    "Team built full-stack radiology AI platform with explainability in 7-day sprint",
    "Domain focus: healthcare AI, medical imaging, clinical decision support",
    "Deliverables: working web app, AI pipeline, Grad-CAM, PDF reports, Docker deployment",
    "Skills: Python/ML, React, FastAPI, DICOM/PACS, system design",
    "Proof: functional prototype in radiology-ai repo with end-to-end demo flow",
]

SLIDE7_BULLETS = [
    "TorchXRayVision — DenseNet121 chest X-ray models: https://github.com/mlmed/torchxrayvision",
    "Grad-CAM explainability — Selvaraju et al.: https://arxiv.org/abs/1610.02391",
    "WHO radiologist workforce data: https://www.who.int/",
    "NIH ChestX-ray14 dataset: https://nihcc.app.box.com/v/ChestXray-NIHCC",
    "Orthanc PACS: https://www.orthanc-server.com/",
    "Captum (PyTorch interpretability): https://captum.ai/",
]


def set_bullets(text_frame, bullets, font_size=16):
    text_frame.clear()
    for idx, bullet in enumerate(bullets):
        para = text_frame.paragraphs[0] if idx == 0 else text_frame.add_paragraph()
        para.text = bullet
        para.level = 0
        para.font.size = Pt(font_size)
        para.font.name = "Arial"


def find_textbox_by_keywords(slide, keywords):
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text.lower()
        if any(k.lower() in text for k in keywords):
            return shape
    return None


def find_empty_content_box(slide):
    candidates = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if shape.name.lower().startswith("textbox"):
            candidates.append(shape)
    if candidates:
        return max(candidates, key=lambda s: s.height * s.width)
    return None


def main():
    prs = Presentation(str(TEMPLATE))

    # Slide 1 — title page metadata
    slide1 = prs.slides[0]
    for shape in slide1.shapes:
        if shape.has_text_frame and "Problem Statement ID" in shape.text:
            shape.text_frame.clear()
            p = shape.text_frame.paragraphs[0]
            p.text = (
                f"Problem Statement ID: {PROBLEM_STATEMENT_ID}\n"
                f"Problem Statement Title: {PROBLEM_STATEMENT_TITLE}\n"
                f"Theme: {THEME}\n"
                f"PS Category: {PS_CATEGORY}"
            )
            for para in shape.text_frame.paragraphs:
                para.font.size = Pt(18)
                para.font.name = "Arial"
                para.font.bold = True

    # Slide 2 — idea title + proposed solution
    slide2 = prs.slides[1]
    for shape in slide2.shapes:
        if shape.has_text_frame and "IDEA TITLE" in shape.text:
            shape.text = IDEA_TITLE
            for para in shape.text_frame.paragraphs:
                para.font.size = Pt(28)
                para.font.bold = True
    content = find_textbox_by_keywords(slide2, ["Proposed Solution", "Detailed explanation"])
    if content:
        tf = content.text_frame
        tf.clear()
        header = tf.paragraphs[0]
        header.text = "Proposed Solution (Describe your Idea / Solution / Prototype)"
        header.font.size = Pt(20)
        header.font.bold = True
        for label in [
            "Detailed explanation of the proposed solution",
            "How it addresses the problem",
            "Innovation and uniqueness of the solution",
        ]:
            p = tf.add_paragraph()
            p.text = label
            p.font.size = Pt(18)
            p.font.bold = True
            p.level = 0
        for bullet in SLIDE2_BULLETS:
            p = tf.add_paragraph()
            p.text = bullet
            p.font.size = Pt(14)
            p.level = 1

    slide_map = {
        2: (SLIDE3_BULLETS, ["TECHNICAL APPROACH", "Technologies to be used"]),
        3: (SLIDE4_BULLETS, ["SOLUTION FEASIBILITY", "Analysis of the feasibility"]),
        4: (SLIDE5_BULLETS, ["REAL WORLD IMPACTS", "Potential impact"]),
        5: (SLIDE6_BULLETS, ["TEAM CAPABILITY", "Team"]),
        6: (SLIDE7_BULLETS, ["RESEARCH", "References"]),
    }

    for slide_idx, (bullets, keywords) in slide_map.items():
        slide = prs.slides[slide_idx]
        box = find_textbox_by_keywords(slide, keywords) or find_empty_content_box(slide)
        if box:
            set_bullets(box.text_frame, bullets, font_size=15)

    # Remove slide 8 (Important Instructions) — optional per template note
    # Keep it so user can review instructions; they delete before upload.

    prs.save(str(OUTPUT))
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    main()
