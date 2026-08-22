"""Update VITISH idea PPTX content to match the current product (keep SIH template pointers)."""

from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

SRC = Path(r"C:\Users\abhin\Downloads\VITISH'26_RadiologyAI_Idea.pptx")
BACKUP = Path(r"C:\Users\abhin\Downloads\VITISH'26_RadiologyAI_Idea_backup.pptx")
OUT_PPTX = Path(r"C:\Users\abhin\Downloads\VITISH'26_RadiologyAI_Idea_UPDATED.pptx")
OUT_PPTX_PROJECT = Path(r"C:\Users\abhin\OneDrive\Desktop\cursor project\radiology-ai\docs\VITISH26_RadiologyAI_Idea.pptx")
OUT_PPTX_SUBMIT = Path(r"C:\Users\abhin\Downloads\VITISH'26_RadiologyAI_Idea_SUBMIT.pptx")


def set_paragraph_text(paragraph, text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.text = text
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def replace_shape_paragraphs(shape, start: int, new_lines: list[str]) -> None:
    paras = shape.text_frame.paragraphs
    for i, line in enumerate(new_lines):
        idx = start + i
        if idx >= len(paras):
            raise IndexError(f"{shape.name}: need paragraph {idx}, have {len(paras)}")
        set_paragraph_text(paras[idx], line)


def delete_slide(prs: Presentation, index: int) -> None:
    sldIdLst = prs.slides._sldIdLst
    sldId = sldIdLst[index]
    rId = sldId.get(qn("r:id"))
    prs.part.drop_rel(rId)
    sldIdLst.remove(sldId)


def main() -> None:
    if not BACKUP.exists():
        BACKUP.write_bytes(SRC.read_bytes())

    prs = Presentation(str(SRC))
    slides = list(prs.slides)

    # Slide 2 — Proposed solution (keep p0–p3 official pointers)
    s2 = slides[1]
    title = next(sh for sh in s2.shapes if sh.name == "Title 1")
    set_paragraph_text(
        title.text_frame.paragraphs[0],
        "Radiology AI Assistant — Explainable Multi-Modality Clinical Decision Support",
    )
    body = next(sh for sh in s2.shapes if sh.name == "TextBox 8")
    replace_shape_paragraphs(
        body,
        4,
        [
            "Web CDS platform for X-ray, CT, MRI, and Ultrasound — radiologist assistant, not a replacement",
            "Upload PNG/JPG/DICOM → modality router; DenseNet121 (CheXNet) predicts 18 chest X-ray findings",
            "CT nodules, MRI tumor volume, US BI-RADS as live pipelines (nnU-Net / MONAI / MedSAM next)",
            "Grad-CAM heatmaps + confidence + severity + ranked differential diagnosis (explainable AI)",
            "Auto radiology report + PDF; FHIR ImagingStudy & DiagnosticReport; Orthanc PACS / DICOMweb",
            "Low-confidence cases flagged ‘Needs Human Review’ — final interpretation by radiologist",
            "Addresses India’s radiologist shortage (~1 per 100,000 vs WHO norm 1:10,000)",
            "Cuts rural reporting delays from 3–7 days to seconds for preliminary triage",
            "Unique: prediction + evidence + report + PACS/FHIR/async worker — full hospital workflow",
        ],
    )

    # Slide 3 — Technologies (keep official pointer textbox)
    s3 = slides[2]
    tech = next(sh for sh in s3.shapes if sh.name == "Title 1")
    replace_shape_paragraphs(
        tech,
        0,
        [
            "AI/ML: PyTorch, DenseNet121 (TorchXRayVision), Grad-CAM (Captum), OpenCV, CLAHE",
            "Backend: FastAPI, SQLAlchemy, SQLite, JWT + RBAC, Celery/Redis async DICOM worker",
            "Frontend: React, Tailwind CSS, Chart.js dashboard, Cornerstone.js DICOM viewer",
            "Imaging: pydicom, Orthanc PACS, DICOMweb QIDO-RS/WADO-RS, HL7 FHIR",
            "Deploy: Docker Compose (API + Nginx + Redis + worker); optional Orthanc --profile pacs",
            "Pipeline: Upload → encrypt → preprocess → modality AI → Grad-CAM → report → dashboard",
            "Working prototype: XR CheXNet + CT/MRI/US demos, audit logs, encrypted images at rest",
        ],
    )

    # Slide 4 — Feasibility
    s4 = slides[3]
    feas = next(sh for sh in s4.shapes if sh.name == "Title 1")
    replace_shape_paragraphs(
        feas,
        0,
        [
            "Feasible: open-source stack, pretrained DenseNet121, Docker on modest hardware",
            "MVP: full-stack XR CDS + live CT/MRI/US routes, PDF reports, PACS/FHIR hooks",
            "Challenges: clinical validation, dataset bias, 3D CT/MRI weights, regulatory path",
            "Risk: false negatives — confidence thresholds, Grad-CAM, radiologist-in-the-loop",
            "Risk: modality mismatch / machine drift — DICOM tags + declared modality + CLAHE",
            "Strategy: hospital pilot on X-ray first; attach nnU-Net / MONAI / MedSAM next",
        ],
    )

    # Slide 5 — Impact
    s5 = slides[4]
    impact = next(sh for sh in s5.shapes if sh.name == "Title 1")
    replace_shape_paragraphs(
        impact,
        0,
        [
            "Rural doctors: instant AI-assisted findings for X-ray/CT/MRI/US without waiting days",
            "Radiologists: draft report ~30 min → ~30 sec; heatmap to verify AI reasoning",
            "Patients: faster pneumonia/TB/nodule triage — 2,500 children die daily from pneumonia",
            "Hospitals: less imaging backlog, PACS/FHIR-ready, lower tele-radiology cost",
            "Social: equitable CDS in underserved areas; RBAC, audit logs, PHI encryption",
            "Economic: India medical imaging ~$2.1B by 2028; tele-radiology growing ~23%/year",
        ],
    )

    # Slide 6 — Team
    s6 = slides[5]
    team = next(sh for sh in s6.shapes if sh.name == "Title 1")
    replace_shape_paragraphs(
        team,
        0,
        [
            "Team built full-stack radiology AI (XR DenseNet + CT/MRI/US + PACS) in hackathon sprint",
            "Domain: healthcare AI, multi-modality imaging, hospital interoperability (PACS/FHIR)",
            "Deliverables: web app, Celery worker, Grad-CAM, PDF reports, FHIR APIs, Docker",
            "Skills: Python/ML, React, FastAPI, DICOM/PACS/FHIR, Docker Compose",
            "Proof: running prototype (docker compose up) — demo doctor@radiology.ai / demo1234",
        ],
    )

    # Slide 7 — References
    s7 = slides[6]
    refs = next(sh for sh in s7.shapes if sh.name == "Title 1")
    replace_shape_paragraphs(
        refs,
        0,
        [
            "TorchXRayVision DenseNet121: https://github.com/mlmed/torchxrayvision",
            "Grad-CAM — Selvaraju et al.: https://arxiv.org/abs/1610.02391 | Captum: https://captum.ai/",
            "NIH ChestX-ray14: https://nihcc.app.box.com/v/ChestXray-NIHCC | WHO workforce: https://www.who.int/",
            "CT LIDC-IDRI + nnU-Net; MRI BraTS + MONAI; US BUSI + MedSAM (planned 3D models)",
            "Orthanc PACS / DICOMweb: https://www.orthanc-server.com/",
            "HL7 FHIR ImagingStudy + DiagnosticReport: https://www.hl7.org/fhir/",
        ],
    )

    prs.save(str(OUT_PPTX))
    OUT_PPTX_PROJECT.parent.mkdir(parents=True, exist_ok=True)
    OUT_PPTX_PROJECT.write_bytes(OUT_PPTX.read_bytes())

    # Portal copy: drop instructions slide 8
    submit = Presentation(str(OUT_PPTX))
    delete_slide(submit, 7)
    submit.save(str(OUT_PPTX_SUBMIT))
    print("saved", OUT_PPTX)
    print("saved", OUT_PPTX_PROJECT)
    print("saved", OUT_PPTX_SUBMIT, "slides", len(submit.slides))


if __name__ == "__main__":
    main()
