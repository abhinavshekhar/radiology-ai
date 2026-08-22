from pptx import Presentation

p = Presentation(r"c:\Users\abhin\Downloads\VITISH'26_RadiologyAI_Idea.pptx")
for i, s in enumerate(p.slides, 1):
    print(f"=== Slide {i} ===")
    for sh in s.shapes:
        if hasattr(sh, "text") and sh.text.strip():
            preview = sh.text.replace("\n", " | ")[:150]
            print(f"  {preview}")
